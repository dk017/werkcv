"""Generates fictional CV files for the CV-check layout tests.

Run from the repo root:  python lib/cv-check/golden/fixtures/generate_fixtures.py
All names and details are fictional.
"""
from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

OUT = Path(__file__).parent
WIDTH, HEIGHT = A4

EXPERIENCE = [
    "Medewerker klantenservice, Voorbeeld Energie BV (2023 - heden)",
    "- Beantwoordde gemiddeld 60 klantvragen per dag",
    "- Verlaagde de afhandeltijd met 15%",
    "- Registreerde klantcontacten in Salesforce",
    "- Begeleidde twee nieuwe collega's bij het inwerken",
    "- Signaleerde terugkerende klachten aan het team",
    "Kassamedewerker, Voorbeeld Supermarkt (2021 - 2022)",
    "- Verwerkte 150 transacties per dienst",
    "- Hielp klanten aan de servicebalie",
    "- Handelde retouren en klachten vriendelijk af",
    "Vakkenvuller, Voorbeeld Supermarkt (2019 - 2021)",
    "- Vulde schappen aan en controleerde houdbaarheid",
    "- Werkte samen in een team van acht collega's",
]


def single_column_pdf() -> None:
    pdf = canvas.Canvas(str(OUT / "single-column.pdf"), pagesize=A4)
    y = HEIGHT - 60
    lines = [
        "Sanne de Vries",
        "Utrecht | sanne.voorbeeld@example.com | 06-12345678",
        "Profiel",
        "Klantgerichte medewerker klantenservice met 3 jaar ervaring in telefonisch",
        "en schriftelijk klantcontact. Lost vragen en klachten snel en vriendelijk op.",
        "Werkervaring",
        *EXPERIENCE,
        "Opleiding",
        "MBO 4 Commercieel medewerker, ROC Midden Nederland (2018 - 2021)",
        "Talen",
        "Nederlands (moedertaal), Engels (C1)",
    ]
    for line in lines:
        pdf.drawString(60, y, line)
        y -= 20
    pdf.save()


def two_column_pdf() -> None:
    pdf = canvas.Canvas(str(OUT / "two-column.pdf"), pagesize=A4)
    left = ["Sanne de Vries", "Werkervaring", *EXPERIENCE, "Opleiding", "MBO 4, ROC (2018 - 2021)"]
    right = ["Contact", "sanne.voorbeeld@example.com", "06-12345678", "Vaardigheden", "Salesforce", "Excel", "Talen", "Nederlands", "Engels"]
    y = HEIGHT - 60
    for index in range(max(len(left), len(right))):
        if index < len(left):
            pdf.drawString(40, y, left[index][:45])
        if index < len(right):
            pdf.drawString(380, y, right[index])
        y -= 20
    pdf.save()


def docx_contact_in_header() -> None:
    document = Document()
    header = document.sections[0].header.paragraphs[0]
    header.text = "sanne.voorbeeld@example.com | 06-12345678"
    document.add_paragraph("Sanne de Vries")
    document.add_paragraph("Werkervaring")
    for line in EXPERIENCE:
        document.add_paragraph(line)
    document.add_paragraph("Opleiding")
    document.add_paragraph("MBO 4 Commercieel medewerker (2018 - 2021)")
    document.save(str(OUT / "contact-in-header.docx"))


if __name__ == "__main__":
    single_column_pdf()
    two_column_pdf()
    docx_contact_in_header()
    print("fixtures written to", OUT)
